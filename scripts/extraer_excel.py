#!/usr/bin/env python3
"""
Extrae los datos introducidos a mano en el Excel de continuidades
(excel/Continuidades_calidad_2024-2030.xlsx) y los guarda en datos/continuidades.json.

Solo lee las ENTRADAS (días, V/B y casillas de abonado), no las fórmulas:
los débitos se recalculan en la app con las reglas de docs/reglas-de-calculo.md.

Uso:  python3 scripts/extraer_excel.py [ruta_excel] [ruta_json]
Requiere: openpyxl
"""
import json, sys, datetime
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "excel" / "Continuidades_calidad_2024-2030.xlsx"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else RAIZ / "datos" / "continuidades.json"

# Cada mes ocupa 6 filas: 5 filas de días (a..a+4) y la fila de subtotal (a+5).
# Rafa: días en C, casilla "Abonado" en B (fila subtotal).
# Carlos: días en I, casilla "Abonado" en H (fila subtotal).
PERSONAS = {"rafa": ("C", "B"), "carlos": ("I", "H")}

def leer():
    wb = openpyxl.load_workbook(XLSX, data_only=False)
    meses = {}
    avisos = []
    for nombre in wb.sheetnames:
        if not nombre.isdigit():
            continue
        ws, anio = wb[nombre], int(nombre)
        for m in range(1, 13):
            a = 6 + 6 * (m - 1)
            sub = a + 5
            clave = f"{anio}-{m:02d}"
            registro = {}
            for p, (col_dias, col_abon) in PERSONAS.items():
                dias, ausencia = [], None
                for r in range(a, a + 5):
                    v = ws[f"{col_dias}{r}"].value
                    if v is None or v == "":
                        continue
                    if isinstance(v, str) and v.strip().upper() in ("V", "B"):
                        ausencia = "vacaciones" if v.strip().upper() == "V" else "baja"
                        continue
                    try:
                        d = int(v)
                        datetime.date(anio, m, d)  # valida la fecha
                        dias.append(d)
                    except (ValueError, TypeError):
                        avisos.append(f"{clave} {p}: valor no válido '{v}' en {col_dias}{r} (se conserva como dia_invalido)")
                        registro.setdefault("_dias_invalidos", {}).setdefault(p, []).append(v)
                abonado = ws[f"{col_abon}{sub}"].value is True
                if dias or ausencia or abonado:
                    registro[p] = {"dias": sorted(dias), "ausencia": ausencia, "abonado": abonado}
            if registro:
                meses[clave] = registro
    return meses, avisos

if __name__ == "__main__":
    meses, avisos = leer()
    datos = {
        "version": 1,
        "generado": datetime.date.today().isoformat(),
        "origen": XLSX.name,
        "meses": dict(sorted(meses.items())),
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Guardado {OUT} ({len(meses)} meses)")
    for a in avisos:
        print("AVISO:", a)
